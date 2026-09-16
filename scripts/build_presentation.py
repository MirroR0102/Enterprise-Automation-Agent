"""Generate Part-4 presentation PPTX + speech PDF + runbook PDF (local only, not for git)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "presentation"
OUT.mkdir(parents=True, exist_ok=True)

PPTX_PATH = OUT / "第4部分-企业业务流程自动化Agent-汇报.pptx"
PDF_PATH = OUT / "第4部分-企业业务流程自动化Agent-演讲稿.pdf"
RUNBOOK_PDF = OUT / "第4部分-企业业务流程自动化Agent-报告执行说明.pdf"
RUNBOOK_MD = OUT / "第4部分-企业业务流程自动化Agent-报告执行说明.md"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0xC4, 0x8A, 0x1A)
INK = RGBColor(0x1E, 0x24, 0x30)
MUTED = RGBColor(0x4A, 0x55, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF6, 0xFA)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
SOFT = RGBColor(0xE8, 0xEE, 0xF8)


def _set_run(run, size=18, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_title_bar(slide, title: str):
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tf = bar.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "  " + title
    _set_run(run, size=22, bold=True, color=WHITE)


def add_footer(slide, page: str):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(12.3), Inches(0.3))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"第 4 部分 · 企业业务流程自动化 Agent　　{page}"
    _set_run(r, size=10, color=MUTED)


def add_bullets(slide, lines: list[str], top=1.15, left=0.7, width=12, size=17):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.7))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        _set_run(run, size=size, color=INK)


def round_box(slide, x, y, w, h, text, fill=LIGHT, edge=NAVY, size=14, bold=True):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = edge
    tf = shp.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    _set_run(r, size=size, bold=bold, color=NAVY)
    return shp


def arrow_right(slide, x, y, w=0.35, h=0.28):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = ACCENT
    shp.line.fill.background()


def caption(slide, x, y, w, text, size=12, color=MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    _set_run(r, size=size, color=color)


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    n = 0

    def next_page():
        nonlocal n
        n += 1
        return f"{n}/20"

    # 1 cover
    s = blank_slide(prs)
    bg = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(1), Inches(2.1), Inches(11.3), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "企业业务流程自动化 Agent"
    _set_run(r, size=36, bold=True, color=WHITE)
    st = s.shapes.add_textbox(Inches(1), Inches(3.4), Inches(11.3), Inches(0.8))
    p = st.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "第 4 部分 · 30 分钟项目汇报"
    _set_run(r, size=22, color=ACCENT)
    meta = s.shapes.add_textbox(Inches(1), Inches(5.1), Inches(11.3), Inches(1))
    p = meta.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "汇报人：________　　日期：________　　MirroR0102/Enterprise-Automation-Agent"
    _set_run(r, size=14, color=WHITE)
    next_page()

    # 2 agenda
    s = blank_slide(prs)
    add_title_bar(s, "目录 · 时间轴（≈30′）")
    items = [
        ("0–3′", "痛点 / 目标 / 第 3↔4 边界"),
        ("3–9′", "架构图 + LangGraph 运行逻辑"),
        ("9–12′", "工具箱 + 安全止损"),
        ("12–20′", "现场演示（可减配）"),
        ("20–26′", "代码地图 / 分库 / 记忆"),
        ("26–30′", "测试、风险、总结、Q&A"),
    ]
    for i, (tmin, desc) in enumerate(items):
        y = 1.2 + i * 0.85
        round_box(s, 0.7, y, 2.2, 0.65, tmin, fill=SOFT, size=16)
        round_box(s, 3.2, y, 9.3, 0.65, desc, fill=LIGHT, size=16, bold=False)
    add_footer(s, next_page())

    # 3 pain + goal
    s = blank_slide(prs)
    add_title_bar(s, "痛点 → 第 4 部分目标")
    round_box(s, 0.6, 1.3, 5.8, 5.2, "", fill=LIGHT)
    caption(s, 0.9, 1.5, 5, "运营痛点", 16, NAVY)
    add_bullets(
        s,
        ["搜新闻 / 查销售 / 算同比 / 写周报步骤碎", "固定脚本跟不上自然语言需求", "数字必须来自工具，过程要可看可停"],
        top=2.1,
        left=0.9,
        width=5.2,
        size=15,
    )
    round_box(s, 6.9, 1.3, 5.8, 5.2, "", fill=SOFT)
    caption(s, 7.2, 1.5, 5, "本部分目标", 16, NAVY)
    add_bullets(
        s,
        ["一句话 → 最多 8 轮工具链 → Markdown 周报", "工作量目标：约减半（需求目标，非已测 KPI）", "JWT ops/dev、时间线、取消、过往周报"],
        top=2.1,
        left=7.2,
        width=5.2,
        size=15,
    )
    add_footer(s, next_page())

    # 4 boundary
    s = blank_slide(prs)
    add_title_bar(s, "大项目边界：第 3 部分 ↔ 第 4 部分")
    round_box(s, 0.7, 2.2, 5.2, 2.8, "第 3 部分\n企业内部知识库 RAG\n向量库 / 上传 / Reranker\n（本仓库不实现）", size=16)
    arrow_right(s, 6.2, 3.3, 0.7, 0.4)
    round_box(
        s,
        7.2,
        2.2,
        5.4,
        2.8,
        "第 4 部分（本交付）\nLangGraph 业务 Agent\n仅预留 enterprise_knowledge_search\n默认 KB Stub，可独立验收",
        size=16,
    )
    caption(s, 0.7, 5.4, 12, "原则：第 4 部分不硬依赖第 3 部分上线；Stub 固定文案「当前未接入企业内部知识库」。", 14)
    add_footer(s, next_page())

    # 5 architecture diagram
    s = blank_slide(prs)
    add_title_bar(s, "架构总览（示意图）")
    layers = [
        (1.2, "浏览器：登录 · 运营工作台 · 过往周报 · 开发日志"),
        (2.3, "FastAPI + JWT · 会话运行时 · 取消标志 · /api/reports"),
        (3.4, "LangGraph StateGraph + MemorySaver（thread_id = session_id）"),
        (4.5, "Tools：Tavily / 计算器 / 时间 / 只读 SQL / 文件 / KB Stub"),
        (5.6, "存储：业务库 sales｜用户分库 eoa_u_{id} 或 data/users/u_{id}.db"),
    ]
    for y, text in layers:
        round_box(s, 0.8, y, 11.7, 0.85, text, fill=LIGHT if int(y * 10) % 2 else SOFT, size=15)
    add_footer(s, next_page())

    # 6 LangGraph runtime
    s = blank_slide(prs)
    add_title_bar(s, "LangGraph 运行逻辑（示意图）")
    round_box(s, 0.5, 2.6, 1.8, 1.0, "START", fill=SOFT, size=14)
    arrow_right(s, 2.4, 2.95)
    round_box(s, 2.9, 2.6, 2.2, 1.0, "agent\n思考/选工具", size=13)
    arrow_right(s, 5.25, 2.95)
    round_box(s, 5.75, 2.6, 2.2, 1.0, "tools\n执行+重试1", size=13)
    arrow_right(s, 8.1, 2.95)
    round_box(s, 8.6, 2.6, 2.0, 1.0, "agent\n再决策", size=13)
    arrow_right(s, 10.75, 2.95)
    round_box(s, 11.2, 2.6, 1.6, 1.0, "finalize", fill=TEAL, size=13)
    caption(s, 0.6, 1.2, 12, "循环条件：有 tool_calls → tools；否则 finalize。止损：轮次>8 / 超时30s / 取消 / 工具连续失败 → END", 13)
    round_box(s, 0.6, 4.3, 3.8, 2.0, "状态字段\nmessages · tool_round\nstatus · events · session_id", size=13)
    round_box(s, 4.7, 4.3, 3.8, 2.0, "status\nrunning / completed\nfailed / cancelled\ntimeout / max_rounds", size=13)
    round_box(s, 8.8, 4.3, 3.9, 2.0, "取消互斥\n已有 final 或终态\n→ 拒绝改写 cancelled", size=13)
    add_footer(s, next_page())

    # 7 toolbox
    s = blank_slide(prs)
    add_title_bar(s, "工具箱（一页看清）")
    tools = [
        (0.5, 1.3, "Tavily 搜索", "公开行业新闻"),
        (3.5, 1.3, "计算器", "ast 白名单表达式"),
        (6.5, 1.3, "时间", "当前日期时间"),
        (9.5, 1.3, "只读 SQL", "sales(sale_date…)"),
        (2.0, 3.6, "文件读写", "仅 reports/"),
        (5.0, 3.6, "KB Stub", "第 3 部分预留"),
        (8.0, 3.6, "用户分库", "周报/会话落库"),
    ]
    for x, y, title, sub in tools:
        round_box(s, x, y, 2.8, 1.7, f"{title}\n{sub}", size=14)
    add_footer(s, next_page())

    # 8 login / session / UI
    s = blank_slide(prs)
    add_title_bar(s, "登录 · 会话 · 界面壳")
    round_box(s, 0.5, 1.3, 4.0, 5.2, "角色\nops：对话/自己的过程\ndev：+ /api/logs\n\n账号\nops/ops123\ndev/dev123", size=15)
    round_box(
        s,
        4.7,
        1.3,
        4.0,
        5.2,
        "会话\nsession_id = thread_id\n同会话可追问\n「新开对话」才清空\n\n事件轮询 / SSE",
        size=15,
    )
    round_box(
        s,
        8.9,
        1.3,
        3.9,
        5.2,
        "界面\n左：能力栏+用户区\n底：composer\n中：过程时间线\n右：周报分屏预览\n过往周报抽屉",
        size=15,
    )
    add_footer(s, next_page())

    # 9 acceptance sequence
    s = blank_slide(prs)
    add_title_bar(s, "验收任务时序（示意图）")
    steps = [
        "用户发验收句",
        "Tavily 搜新闻",
        "SQL 查本月销售",
        "计算器同比",
        "写 Markdown",
        "时间线 + 周报",
    ]
    for i, text in enumerate(steps):
        x = 0.4 + i * 2.15
        round_box(s, x, 2.8, 2.0, 1.4, f"{i + 1}\n{text}", size=13)
        if i < len(steps) - 1:
            arrow_right(s, x + 2.05, 3.35, 0.28, 0.25)
    caption(
        s,
        0.5,
        4.8,
        12,
        "种子数据：本月 100000 / 去年同月 80000 → 同比 25%。SQL 列名必须是 sale_date，不是 order_date。",
        13,
    )
    add_footer(s, next_page())

    # 10 demo shot placeholders
    s = blank_slide(prs)
    add_title_bar(s, "演示画面位（现场替换截图）")
    round_box(s, 0.5, 1.3, 6.0, 5.2, "【截图位 A】\n工作台：侧栏 + 时间线 + composer\n指：用户区 / 过程 / 状态芯片", size=16)
    round_box(s, 6.8, 1.3, 6.0, 5.2, "【截图位 B】\n周报分屏预览 / 过往周报抽屉\n指：Markdown 与用户库落库", size=16)
    add_footer(s, next_page())

    # 11 memory demo
    s = blank_slide(prs)
    add_title_bar(s, "多轮记忆演示（同 session）")
    round_box(s, 0.7, 2.0, 5.5, 3.5, "第 1 轮\n请按「上周格式」生成周报\n→ 产出结构被记住", size=16)
    arrow_right(s, 6.4, 3.5, 0.6, 0.35)
    round_box(s, 7.2, 2.0, 5.4, 3.5, "第 2 轮（同一会话）\n本周数据请沿用同一格式\n→ MemorySaver 续写\n点「新开对话」才重置", size=16)
    add_footer(s, next_page())

    # 12 code map
    s = blank_slide(prs)
    add_title_bar(s, "代码地图（评委爱问「在哪改」）")
    rows = [
        ("图 / 轮次 / 路由", "app/agent/graph.py · nodes.py · state.py"),
        ("系统提示", "app/agent/prompts.py"),
        ("工具注册与 SQL 防护", "app/tools/registry.py · mysql_query.py"),
        ("会话 / 取消 / 周报落库", "app/api/chat.py · reports.py · runtime.py"),
        ("用户分库", "app/db/user_store.py"),
        ("前端壳", "web/index.html · app.css · common.js"),
    ]
    for i, (k, v) in enumerate(rows):
        y = 1.15 + i * 0.9
        round_box(s, 0.5, y, 3.8, 0.75, k, fill=SOFT, size=14)
        round_box(s, 4.5, y, 8.3, 0.75, v, fill=LIGHT, size=14, bold=False)
    add_footer(s, next_page())

    # 13 per-user store
    s = blank_slide(prs)
    add_title_bar(s, "每用户分库 + 过往周报（已交付）")
    round_box(s, 0.6, 1.5, 6.0, 4.8, "业务库\nenterprise_ops / mock.db\nsales · users · agent_logs\n→ 只读查询工具", size=16)
    round_box(
        s,
        6.9,
        1.5,
        5.9,
        4.8,
        "用户库\nMySQL: eoa_u_{id}\nmock: data/users/u_{id}.db\n会话 / 消息 / weekly_reports\nAPI: GET /api/reports",
        size=16,
    )
    add_footer(s, next_page())

    # 14 safety
    s = blank_slide(prs)
    add_title_bar(s, "安全与止损图")
    fears = [
        ("死循环", "最大工具轮次 8"),
        ("挂太久", "TASK_TIMEOUT_S=30"),
        ("外网抖", "失败重试 1 次再停"),
        ("SQL 误伤", "仅 SELECT + 黑名单"),
        ("路径穿越", "文件仅 reports/"),
        ("误取消", "终态/final 不改写"),
    ]
    for i, (a, b) in enumerate(fears):
        col, row = i % 3, i // 3
        x, y = 0.6 + col * 4.2, 1.4 + row * 2.5
        round_box(s, x, y, 3.9, 2.1, f"怕：{a}\n→ {b}", size=15)
    add_footer(s, next_page())

    # 15 tests
    s = blank_slide(prs)
    add_title_bar(s, "测试与验收口径")
    add_bullets(
        s,
        [
            "pytest：鉴权、取消互斥、SQL 防护、计算器、KB Stub、图轮次、过往周报隔离…",
            "联网验收：DeepSeek + Tavily 配齐后跑 scripts/acceptance_check.md",
            "演示账号 ops/ops123、dev/dev123；本地 http://127.0.0.1:8000",
            "联调结论见 README「联调问题记录」（order_date → 错误回传自纠）",
        ],
        size=16,
    )
    add_footer(s, next_page())

    # 16 risks
    s = blank_slide(prs)
    add_title_bar(s, "风险与诚实边界")
    add_bullets(
        s,
        [
            "MemorySaver 进程内：多副本/重启丢图状态（日志仍在库）",
            "30s 超时对长链路偏紧：演示可临时调配置，规划默认不放宽",
            "Tavily/DeepSeek 未配置时必须明示失败，禁止编造新闻与销售额",
            "不做：OAuth、细粒度 RBAC、K8s、真实数仓、第 3 部分 RAG 本体",
        ],
        size=16,
    )
    add_footer(s, next_page())

    # 17 summary
    s = blank_slide(prs)
    add_title_bar(s, "总结：交付了什么")
    add_bullets(
        s,
        [
            "可运行的第 4 部分：LangGraph Agent + FastAPI + 浅色工作台",
            "安全止损齐全；SQL 自纠；取消与终态互斥",
            "同会话记忆、昵称设置、每用户分库与过往周报",
            "文档：README / AGENT / 规划 / 本汇报三件套（本地）",
        ],
        size=17,
    )
    add_footer(s, next_page())

    # 18 outlook
    s = blank_slide(prs)
    add_title_bar(s, "展望")
    add_bullets(
        s,
        [
            "接入真实第 3 部分 KB（关掉 Stub）",
            "Postgres checkpointer 替换 MemorySaver",
            "更细的会话列表与周报模板管理",
            "演示环境密钥与超时策略的运维手册化",
        ],
        size=17,
    )
    add_footer(s, next_page())

    # 19 Q&A
    s = blank_slide(prs)
    add_title_bar(s, "Q & A")
    round_box(
        s,
        2.5,
        2.5,
        8.3,
        2.5,
        "欢迎提问\n备答：为何 LangGraph / 8 轮 / 30s / KB Stub / 分库隔离",
        size=18,
    )
    add_footer(s, next_page())

    # 20 thank you
    s = blank_slide(prs)
    bg = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.3), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "谢谢 · 请批评指正"
    _set_run(r, size=36, bold=True, color=WHITE)
    st = s.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.3), Inches(0.8))
    p = st.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "PPT / 讲稿 / 执行说明 · 本地 docs/presentation · 不推送 GitHub"
    _set_run(r, size=16, color=ACCENT)
    next_page()

    prs.save(str(PPTX_PATH))
    return PPTX_PATH


def _register_font() -> str:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\msyh.ttf"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    for path in candidates:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont("CN", str(path)))
                return "CN"
            except Exception:
                continue
    return "Helvetica"


def _pdf_styles(font: str):
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "CNTitle",
        parent=styles["Title"],
        fontName=font,
        fontSize=16,
        leading=22,
        spaceAfter=12,
        textColor=HexColor("#1B2A4A"),
    )
    h1 = ParagraphStyle(
        "CNH1",
        parent=styles["Heading1"],
        fontName=font,
        fontSize=13,
        leading=18,
        spaceBefore=12,
        spaceAfter=6,
        textColor=HexColor("#1B2A4A"),
    )
    body = ParagraphStyle(
        "CNBody",
        parent=styles["Normal"],
        fontName=font,
        fontSize=10.5,
        leading=16,
        spaceAfter=6,
        textColor=HexColor("#1E2430"),
    )
    cue = ParagraphStyle(
        "CNCue",
        parent=body,
        textColor=HexColor("#8B1E1E"),
        spaceBefore=4,
        spaceAfter=8,
    )
    return title, h1, body, cue


def build_speech_pdf() -> Path:
    font = _register_font()
    title, h1, body, cue = _pdf_styles(font)
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
    )
    story: list = []

    def P(text: str, style=body):
        story.append(Paragraph(text.replace("\n", "<br/>"), style))

    def B(text: str):
        P(f"<b>{text}</b>", cue)

    P("第 4 部分 · 企业业务流程自动化 Agent · 演讲稿（约 30 分钟口语完整稿）", title)
    P(
        "配套 PPT：同目录「汇报.pptx」共 20 页。时间轴：开场 3′ · 架构+运行 6′ · 工具安全 3′ · 演示 8′ · 代码/分库 6′ · 总结 Q&A 4′。"
        "【动作】段落只给汇报人看，不要念出声。减配：砍展望/风险细讲，绝不砍演示。加餐：多走一轮记忆追问或打开过往周报。",
        body,
    )
    P("汇报人：________　　日期：________　　仓库：MirroR0102/Enterprise-Automation-Agent", body)

    P("0:00–3:00 · 开场 / 痛点 / 边界（PPT 1–4）", h1)
    B("【动作】打开 PPT 封面，报题目与仓库名；不要念文件路径。")
    P(
        "各位老师、同学好。今天汇报大项目第 4 部分——企业业务流程自动化 Agent。"
        "运营同学写周报要搜新闻、查销售额、算同比、拼 Markdown，步骤碎、脚本一改需求就废。"
        "我们希望他说一句话，系统用有状态 Agent 规划并调用工具，过程可见、可取消，数字来自工具而不是模型瞎编。"
        "业务目标写的是周报类工作量约减半——这是需求目标，不是我们已经测过的 KPI。"
        "大项目里第 3 部分是知识库 RAG；我们不实现向量库，只预留 enterprise_knowledge_search，默认 Stub，"
        "所以第 4 部分可以独立验收，不会被第 3 部分拖死。",
        body,
    )

    P("3:00–9:00 · 架构 + LangGraph 运行逻辑（PPT 5–6）", h1)
    B("【动作】指架构五层盒子，再指运行逻辑从 START 到 finalize 的箭头。")
    P(
        "架构从上到下：浏览器工作台；FastAPI 加 JWT 和取消标志；LangGraph StateGraph 加 MemorySaver，"
        "thread_id 等于 session_id；下面是工具；最底是业务库和每用户分库。"
        "为什么用图而不是一次性 Prompt？因为要多步工具、条件结束、控轮次和取消。"
        "节点很简单：agent 思考并可能发起 tool_calls，有调用就进 tools，tools 回来再进 agent；"
        "没有工具调用就 finalize。止损有四条：工具轮次超过 8、任务超时默认 30 秒、用户取消、工具连续失败。"
        "状态里有 messages、tool_round、status、events。特别说明：任务已经 completed 或者事件里已有 final，"
        "迟到的取消不能把状态改成 cancelled——评委如果问「做完了还能取消吗」，答案是不能误标。",
        body,
    )

    P("9:00–12:00 · 工具箱与安全止损（PPT 7、14）", h1)
    P(
        "工具：Tavily 搜公开信息；计算器做同比；时间戳；只读 SQL 查 sales，列是 sale_date、amount、region；"
        "文件只许落在 reports；知识库默认 Stub；用户分库存会话和周报。"
        "安全用「怕什么挡什么」：怕死循环就 8 轮；怕挂太久就 30 秒；怕外网抖就重试一次再停；"
        "怕 SQL 误伤就只允许 select 加黑名单；怕路径穿越就锁目录；怕误取消就终态互斥。"
        "联调真实踩坑：模型把 sale_date 写成 order_date，若把执行失败当致命错误，整条任务当场死掉；"
        "我们改成错误回传让 Agent 自纠，并在提示词写清 schema。细节在 README 联调问题记录，写报告可直接引用。",
        body,
    )

    P("12:00–20:00 · 现场演示（PPT 8–11）", h1)
    B("【动作】切到浏览器 http://127.0.0.1:8000 ；用 ops/ops123 登录；确认左下角显示 ops。")
    P(
        "先指界面：左侧能力栏和用户区，底部输入框，中间执行过程，出周报后右侧分屏预览。"
        "粘贴验收句：帮我搜索2026AI行业新闻，再查询本月销售总额，计算同比增长率，生成一份运营周报markdown。"
        "发送后指时间线：搜索、SQL、计算器；强调销售额来自数据库。种子数据本月十万、去年八万，同比百分之二十五。"
        "可选：点开过往周报；或同会话追问「请沿用刚才的周报格式」展示记忆；或用 dev 看日志。"
        "若现场超时：立刻切备用录屏或已有 reports 文件，口头说明「按减配保留演示结论」。",
        body,
    )
    B("演示口令：「这里调用了数据库而不是编造销售额。」")
    B("演示口令：「同一会话可以追问；只有新开对话才清空记忆。」")

    P("20:00–26:00 · 代码如何实现（PPT 12–13）≈5–6 分钟", h1)
    B("【动作】可打开仓库树或 PPT 代码地图，不要大段念代码。")
    P(
        "改图结构看 app/agent/graph.py：START 到 agent，条件边进 tools 或 finalize。"
        "轮次和失败重试在 nodes；系统提示在 prompts；工具清单在 registry；SQL 防护在 mysql_query。"
        "API 层 chat.py 负责建会话、投递消息、取消；成功周报会进 user_store 的 weekly_reports。"
        "分库：MySQL 用 eoa_u_{用户id}，mock 用 data/users/u_{id}.db，和销售业务表分离——"
        "查询销售仍走业务库，用户上下文与过往周报走用户库。前端是静态页，不引入重框架，方便课堂演示。",
        body,
    )

    P("26:00–30:00 · 测试、风险、总结、Q&A（PPT 15–20）", h1)
    P(
        "测试用 pytest 覆盖鉴权、取消、SQL、图轮次、周报隔离；联网验收按 acceptance_check。"
        "风险要诚实：MemorySaver 重启丢图状态；30 秒对长链路紧；没配密钥就不能装成功。"
        "交付是可运行代码、测试、演示账号、文档，以及本地三份汇报材料——PPT、讲稿、执行说明，不进 GitHub 推送。"
        "我的汇报到这里，谢谢大家，欢迎提问。",
        body,
    )
    B("备答提示：为何不用纯 AgentExecutor；KB 为何 Stub；分库为何不拆 sales。")

    doc.build(story)
    return PDF_PATH


def build_runbook() -> tuple[Path, Path]:
    font = _register_font()
    title, h1, body, cue = _pdf_styles(font)

    md = """# 第 4 部分 · 报告执行说明（自用，不要念）

> 封面提示：**本文件给汇报人控场，不对听众朗读。**

## 会前清单（T-30′）

- [ ] `.venv` 可用；`uvicorn` 已起 `127.0.0.1:8000`
- [ ] `.env`：DeepSeek / Tavily（演示要联网搜索时）
- [ ] 浏览器无痕或已登录 ops；硬刷新一次
- [ ] PPT / 讲稿 / 本说明三份打开；验收句复制到剪贴板
- [ ] 备用：录屏或 `reports/` 下已有周报；VPN 仅 push 时需要

## 时间轴总表（对齐 PPT 页码）

| 分钟 | PPT | 手上干什么 | 嘴里讲什么 |
| --- | --- | --- | --- |
| 0–3 | 1–4 | 翻页，勿开演示 | 痛点、目标、3/4 边界 |
| 3–9 | 5–6 | 指架构与运行箭头 | LangGraph 循环与止损 |
| 9–12 | 7、14 | 指工具与安全六格 | 工具箱 + 怕什么挡什么 |
| 12–20 | 8–11 | **切浏览器演示** | 验收句、时间线、同比 25% |
| 20–26 | 12–13 | 可切 IDE/代码地图 | 关键文件与分库 |
| 26–30 | 15–20 | 回 PPT | 测试、风险、总结、Q&A |

## 演示原句

```
帮我搜索2026AI行业新闻，再查询本月销售总额，计算同比增长率，生成一份运营周报markdown
```

记忆加餐（同会话）：

```
请沿用刚才周报的标题层级和章节结构，补充一句本周风险提示
```

账号：`ops/ops123`（主演示），`dev/dev123`（日志）。

## 减配 / 加餐

- **减配（超时）**：跳过展望与风险细讲；演示改录屏；代码段缩到 3 分钟只讲 graph.py + user_store。
- **加餐（富余）**：过往周报抽屉；同会话记忆第二轮；点用户区改昵称；dev 日志页。

## 备用方案

1. API/模型失败 → 打开已生成 `reports/*.md`，说明工具链曾跑通。
2. 搜索失败 → 强调「失败明文返回、不编造新闻」，仍可展示 SQL+计算器。
3. 前端身份异常 → Ctrl+F5；确认 `common.js` 已加载。

## 不要做

- 不要照搬合同审查范本业务内容
- 不要把未实现能力（真实 RAG、K8s）说成已交付
- 未开 VPN 不要 push；PPT/PDF 默认不进远程仓库
"""
    RUNBOOK_MD.write_text(md, encoding="utf-8")

    doc = SimpleDocTemplate(
        str(RUNBOOK_PDF),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
    )
    story: list = []

    def P(text: str, style=body):
        story.append(Paragraph(text.replace("\n", "<br/>"), style))

    def B(text: str):
        P(f"<b>{text}</b>", cue)

    P("第 4 部分 · 企业业务流程自动化 Agent · 报告执行说明", title)
    B("本文件给汇报人控场使用，封面起就不要对听众朗读。")
    P("配套：汇报.pptx（20 页）· 演讲稿.pdf · 本说明。仓库：MirroR0102/Enterprise-Automation-Agent", body)

    P("一、会前清单（T-30′）", h1)
    P(
        "确认 venv 与 uvicorn 监听 127.0.0.1:8000；按需配置 DeepSeek/Tavily；浏览器硬刷新；"
        "三份材料打开；验收句在剪贴板；备用录屏或 reports 已有文件。Push 前再开 VPN。",
        body,
    )

    P("二、时间轴总表（与 PPT 页码对齐）", h1)
    P("0–3′ PPT1–4 开场边界｜3–9′ PPT5–6 架构运行｜9–12′ PPT7/14 工具安全｜"
      "12–20′ PPT8–11 现场演示｜20–26′ PPT12–13 代码分库｜26–30′ PPT15–20 总结问答。", body)

    P("三、逐段手上动作", h1)
    P("开场只翻 PPT。演示段必须切浏览器，登录 ops/ops123，发送验收句，指时间线与同比 25%。"
      "代码段指代码地图或 IDE，不念大段源码。收尾回 PPT 总结并留 Q&A。", body)

    P("四、演示原句与加餐", h1)
    P("验收句：帮我搜索2026AI行业新闻，再查询本月销售总额，计算同比增长率，生成一份运营周报markdown", body)
    P("记忆加餐：请沿用刚才周报的标题层级和章节结构，补充一句本周风险提示", body)

    P("五、减配 / 加餐 / 备用", h1)
    P("减配：砍展望细讲、改录屏、代码缩到 graph+user_store。加餐：过往周报、第二轮记忆、改昵称、dev 日志。"
      "备用：模型失败展示已有 Markdown；搜索失败强调不编造；身份异常 Ctrl+F5。", body)

    P("六、禁止事项", h1)
    P("禁止照搬合同范本业务；禁止把 RAG/K8s 说成已交付；禁止未开 VPN 强行 push；汇报二进制默认不进 GitHub。", body)

    doc.build(story)
    return RUNBOOK_PDF, RUNBOOK_MD


def main():
    pptx = build_pptx()
    pdf = build_speech_pdf()
    runbook_pdf, runbook_md = build_runbook()
    print(f"wrote {pptx}")
    print(f"wrote {pdf}")
    print(f"wrote {runbook_pdf}")
    print(f"wrote {runbook_md}")


if __name__ == "__main__":
    main()
